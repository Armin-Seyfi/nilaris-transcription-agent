import copy
import hashlib
import json
import sys
import tempfile
import threading
import time
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
import meeting_notes as notes
import session_titles as titles
from notes_exporter import create_notes_docx

SID='nilaris_2026-10-13T12-00-00Z'
def example(language='en'):
    return {'language':language,'title':'Procurement planning','content_type':'Meeting',
        'overview':'The team compared supplier lead times. A follow-up was agreed.',
        'topics':[{'heading':'Supplier lead times','points':['Availability remains uncertain.']}],
        'decisions':['Request revised lead times.'],
        'actions':[{'action':'Request revised lead times','owner':'','timeline':'next week','context':'For planning'}],
        'open_questions':['Which supplier can confirm delivery?'],'highlights':[], 'uncertainties':[]}

class NotesTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.outputs=self.root/'outputs';self.records=self.root/'recordings'
        self.folder=self.outputs/(SID+'_separated');self.folder.mkdir(parents=True);self.records.mkdir()
        for module,key,value in [(notes,'OUTPUTS_DIR',self.outputs),(titles,'OUTPUTS_DIR',self.outputs),(titles,'RECORDINGS_DIR',self.records)]:
            mock=patch.object(module,key,value);mock.start();self.addCleanup(mock.stop)
        self.source=self.folder/'transcript.json'
        self.source.write_text(json.dumps({'segments':[{'text':'Compare lead times. Follow up next week.','source':'REMOTE','start':0,'end':10}]}))
        (self.folder/'transcript.txt').write_text('Compare lead times. Follow up next week.')
        self.fingerprint=hashlib.sha256(self.source.read_bytes()).hexdigest()

    def wait(self,manager):
        deadline=time.monotonic()+4
        while manager.busy() and time.monotonic()<deadline: time.sleep(.01)
        self.assertFalse(manager.busy())

    def test_saved_word_reused_after_restart_without_calls(self):
        calls=[]
        manager=notes.NotesManager(request_fn=lambda task,payload:(calls.append(task) or example()))
        manager.start(SID);self.wait(manager)
        self.assertEqual(manager.status(SID)['state'],'ready')
        new=notes.NotesManager(request_fn=lambda *args:self.fail('Download or restart must never send AI requests'))
        self.assertEqual(new.start(SID)['state'],'ready')
        self.assertEqual(len(calls),1)
        doc=Document(self.folder/'ai_notes.docx')
        self.assertEqual(len(doc.tables),1)
        self.assertIn('Not specified',doc.tables[0].cell(1,1).text)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(),self.fingerprint)

    def test_double_click_single_request_and_concurrent_session_blocked(self):
        entered=threading.Event();release=threading.Event();calls=[]
        def request(*args): entered.set();release.wait(2);calls.append(1);return example()
        manager=notes.NotesManager(request_fn=request)
        manager.start(SID);self.assertTrue(entered.wait(1));manager.start(SID)
        release.set();self.wait(manager);self.assertEqual(len(calls),1)

    def test_failure_does_not_auto_retry_and_keeps_source(self):
        calls=[]
        def request(*args): calls.append(1);raise RuntimeError('Incomplete notes')
        manager=notes.NotesManager(request_fn=request);manager.start(SID);self.wait(manager)
        for _ in range(3):self.assertEqual(manager.status(SID)['state'],'error')
        self.assertEqual(len(calls),1);self.assertFalse((self.folder/'ai_notes.docx').exists())
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(),self.fingerprint)

    def test_full_long_transcript_and_checkpointed_retry(self):
        self.source.write_text(json.dumps({'segments':[{'text':'opening '+('x'*42000)+' final decision', 'source':'REMOTE'}]}))
        chunks=notes.transcript_chunks(self.folder)[1]
        self.assertGreater(len(chunks),1);self.assertIn('final decision',chunks[-1])
        calls=[]
        def request(task,payload):
            calls.append(task)
            if task.startswith('Synthesise'):raise RuntimeError('synthesis failed')
            return example('de')
        manager=notes.NotesManager(request_fn=request);manager.start(SID);self.wait(manager)
        self.assertEqual(manager.status(SID)['state'],'error')
        manager.request_fn=lambda task,payload:(calls.append(task) or example('de'))
        manager.start(SID);self.wait(manager)
        self.assertEqual(manager.status(SID)['state'],'ready')
        self.assertEqual(len(calls),len(chunks)+2)

    def test_source_changed_during_generation_prevents_publish(self):
        def request(*args):self.source.write_text('{"segments":[]}');return example()
        manager=notes.NotesManager(request_fn=request);manager.start(SID);self.wait(manager)
        self.assertEqual(manager.status(SID)['state'],'error')
        self.assertFalse((self.folder/'ai_notes.docx').exists())

    def test_word_failure_retry_reuses_completed_ai_result(self):
        calls=[]
        manager=notes.NotesManager(request_fn=lambda *args:(calls.append(1) or example()))
        with patch('notes_exporter.create_notes_docx',side_effect=RuntimeError('Word failed')):
            manager.start(SID);self.wait(manager)
        manager.start(SID);self.wait(manager)
        self.assertEqual(manager.status(SID)['state'],'ready');self.assertEqual(len(calls),1)

    def test_safety_and_missing_transcript_no_api(self):
        manager=notes.NotesManager(request_fn=lambda *args:self.fail('Invalid source sent'))
        with self.assertRaises(ValueError):manager.start('../outside')
        (self.folder/'ai_notes.docx').symlink_to(self.root/'external')
        with self.assertRaises(ValueError):manager.start(SID)
        (self.folder/'ai_notes.docx').unlink();self.source.unlink()
        with self.assertRaises(FileNotFoundError):manager.start(SID)

    def test_language_and_word_directions_and_optional_sections(self):
        for language in ('en','de','fa'):
            data=example(language);data['actions']=[];data['decisions']=[]
            if language=='fa':data.update(title='برنامه‌ریزی خرید',overview='تیم زمان تحویل تأمین‌کنندگان را بررسی کرد.')
            doc=Document(BytesIO(create_notes_docx(data,{'session_id':SID})))
            self.assertEqual(len(doc.tables),0)
            self.assertEqual(doc.paragraphs[0].style.name,'Title')
            val=doc.paragraphs[0]._p.pPr.find(qn('w:bidi')).get(qn('w:val'))
            self.assertEqual(val,'1' if language=='fa' else '0')
            if language=='de':self.assertIn('Datum: 13.10.2026',[p.text for p in doc.paragraphs])
        bad=example();bad['actions'][0]['owner']=None
        with self.assertRaises(ValueError):notes.validate_notes(bad)

    def test_word_alignment_follows_microsoft_bidi_semantics_in_every_paragraph(self):
        for language in ('fa', 'en', 'de'):
            doc = Document(BytesIO(create_notes_docx(example(language), {'session_id': SID})))
            paragraphs = [p for p in doc.paragraphs if p.text]
            paragraphs += [p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs if p.text]
            for paragraph in paragraphs:
                self.assertEqual(paragraph._p.pPr.find(qn('w:jc')).get(qn('w:val')), 'left')
                self.assertEqual(paragraph._p.pPr.find(qn('w:bidi')).get(qn('w:val')), '1' if language == 'fa' else '0')

    def test_legacy_persian_word_download_repairs_once_without_ai_and_preserves_backup(self):
        data = example('fa')
        record = {'notes': data, 'model': 'gpt-4o-mini', 'created_at': 'baseline'}
        (self.folder/'ai_notes.json').write_text(json.dumps(record))
        old = create_notes_docx(data, {'session_id': SID})
        (self.folder/'ai_notes.docx').write_bytes(old)
        manager = notes.NotesManager(request_fn=lambda *args: self.fail('Formatting must not call AI'))
        with patch('notes_exporter.create_notes_docx', wraps=create_notes_docx) as render:
            manager.download_document(SID)
            manager.download_document(SID)
            self.assertEqual(render.call_count, 1)
        self.assertEqual((self.folder/'ai_notes.before-alignment.docx').read_bytes(), old)
        saved = json.loads((self.folder/'ai_notes.json').read_text())
        self.assertEqual(saved['notes'], data)
        self.assertEqual(saved['created_at'], 'baseline')
        self.assertEqual(saved['document_format_version'], 2)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.fingerprint)

    def test_legacy_english_and_german_downloads_reuse_original_files(self):
        for language in ('en', 'de'):
            (self.folder/'ai_notes.json').write_text(json.dumps({'notes': example(language)}))
            (self.folder/'ai_notes.docx').write_bytes(b'original cached Word')
            manager = notes.NotesManager(request_fn=lambda *args: self.fail('No AI'))
            with patch('notes_exporter.create_notes_docx') as render:
                self.assertEqual(manager.download_document(SID).read_bytes(), b'original cached Word')
                render.assert_not_called()

    def test_origin_download_and_storage_api(self):
        import web_server as web
        manager=notes.NotesManager(request_fn=lambda *args:example())
        with patch.object(web,'notes_manager',manager):
            client=web.app.test_client()
            url=f'/api/sessions/{SID}'
            self.assertEqual(client.post(url+'/notes/generate').status_code,403)
            self.assertEqual(client.get(url+'/notes/generate').status_code,405)
            self.assertEqual(client.get(url+'/download/notes').status_code,409)
            self.assertEqual(client.post(url+'/notes/generate',headers={'Origin':web.ALLOWED_ORIGIN}).status_code,202)
            self.wait(manager)
            response=client.get(url+'/download/notes')
            self.assertEqual(response.status_code,200);self.assertTrue(response.data.startswith(b'PK'))
            self.assertEqual(client.get(url+'/download/docx').status_code,404)
            with patch.object(manager,'busy',return_value=True):
                with self.assertRaises(RuntimeError):web.ensure_session_storage_available()

if __name__=='__main__':unittest.main()
