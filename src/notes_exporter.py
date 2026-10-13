"""Localized Word formatting reusing existing transcript direction/date helpers."""
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from transcript_exporter import add_paragraph, set_text_direction, get_start_datetime, format_dates, format_duration

LABELS = {
 'en': ['AI Generated Notes', 'Content type', 'Overview', 'Key topics', 'Decisions', 'Actions and plans', 'Open questions', 'Highlights and ideas', 'Uncertainties', 'Action or plan', 'Owner', 'Timeline', 'Context', 'Not specified', 'AI generated from the saved transcript. Review important decisions and actions against the original recording.'],
 'fa': ['یادداشت‌های تولیدشده با هوش مصنوعی', 'نوع محتوا', 'خلاصه', 'موضوعات اصلی', 'تصمیم‌ها', 'اقدام‌ها و برنامه‌ها', 'پرسش‌های باز', 'نکات و ایده‌های برجسته', 'ابهام‌ها', 'اقدام یا برنامه', 'مسئول', 'زمان‌بندی', 'توضیحات', 'مشخص نشده', 'این یادداشت‌ها با هوش مصنوعی از رونوشت ذخیره‌شده تولید شده‌اند. تصمیم‌ها و اقدام‌های مهم را با ضبط اصلی بررسی کنید.'],
 'de': ['KI generierte Notizen', 'Art des Inhalts', 'Überblick', 'Wichtige Themen', 'Entscheidungen', 'Aufgaben und Pläne', 'Offene Fragen', 'Höhepunkte und Ideen', 'Unklarheiten', 'Aufgabe oder Plan', 'Verantwortlich', 'Zeitplan', 'Kontext', 'Nicht angegeben', 'Diese Notizen wurden mit KI aus dem gespeicherten Transkript erstellt. Prüfen Sie wichtige Entscheidungen und Aufgaben anhand der Originalaufnahme.']}


def direction(paragraph, rtl):
    set_text_direction(paragraph, rtl)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        run.font.name = 'Arial'
        run.font.color.rgb = RGBColor(0,0,0)
        fonts = run._r.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is not None: fonts.set(qn('w:cs'), 'Arial')
        if rtl:
            prop = OxmlElement('w:rtl'); prop.set(qn('w:val'),'1'); run._r.get_or_add_rPr().append(prop)


def create_notes_docx(notes, metadata):
    language = notes['language']; labels = LABELS[language]; rtl = language == 'fa'
    document = Document(); section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.8)
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style = document.styles[name]
        style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt(11 if name=='Normal' else 18 if name=='Title' else 13 if name=='Heading 1' else 11)
    document.styles['Normal'].paragraph_format.line_spacing=1.15
    def para(text, style=None):
        p=add_paragraph(document,text,rtl=rtl)
        if style:
            p.style=style
            for run in p.runs:
                run.font.size=Pt(18 if style=='Title' else 13 if style=='Heading 1' else 11)
                run.bold=True
            p.paragraph_format.keep_with_next=True
        direction(p,rtl)
        return p
    para(notes['title'],'Title')
    para(labels[0])
    timestamp=get_start_datetime(metadata,metadata['session_id'])
    dates = format_dates(timestamp,language) if language in ('fa','en') else [('Datum', timestamp.strftime('%d.%m.%Y'))]
    for label,value in dates: para(f'{label}: {value}')
    time_label={'fa':'زمان ضبط به وقت برلین','en':'Recorded at Berlin time','de':'Aufnahmezeit Berlin'}[language]
    para(f'{time_label}: {timestamp.strftime("%H:%M")}')
    if metadata.get('duration_seconds') is not None:
        duration_label={'fa':'مدت ضبط','en':'Recording duration','de':'Aufnahmedauer'}[language]
        para(f'{duration_label}: {format_duration(metadata["duration_seconds"])}')
    if notes['content_type']: para(f'{labels[1]}: {notes["content_type"]}')
    para(labels[2],'Heading 1'); para(notes['overview'])
    if notes['topics']:
        para(labels[3],'Heading 1')
        for topic in notes['topics']:
            para(topic['heading'],'Heading 2')
            for point in topic['points']: para('• '+point)
    if notes['decisions']:
        para(labels[4],'Heading 1')
        for item in notes['decisions']: para('• '+item)
    if notes['actions']:
        para(labels[5],'Heading 1')
        table=document.add_table(rows=1,cols=4);table.autofit=False
        widths=[2.2,1.05,1.05,2.37]
        if rtl:
            prop=OxmlElement('w:bidiVisual');prop.set(qn('w:val'),'1');table._tbl.tblPr.append(prop)
        borders=OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
        table._tbl.tblPr.append(borders)
        for col,width in zip(table.columns,widths): col.width=Inches(width)
        header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
        rows=[labels[9:13]]+[[item['action'],item['owner'] or labels[13],item['timeline'] or labels[13],item['context']] for item in notes['actions']]
        for i,values in enumerate(rows):
            row=table.rows[0] if i==0 else table.add_row()
            for cell,value,width in zip(row.cells,values,widths):
                cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text=value
                for p in cell.paragraphs:
                    direction(p,rtl);p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
                    for run in p.runs: run.font.size=Pt(10);run.bold=(i==0)
                pr=cell._tc.get_or_add_tcPr()
                margins=OxmlElement('w:tcMar')
                for edge in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');margins.append(el)
                pr.append(margins)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF2');pr.append(shade)
        document.add_paragraph()
    for key,index in [('open_questions',6),('highlights',7),('uncertainties',8)]:
        if notes[key]:
            para(labels[index],'Heading 1')
            for item in notes[key]: para('• '+item)
    p=para(labels[14]);p.paragraph_format.space_before=Pt(14)
    for run in p.runs: run.font.size=Pt(9)
    output=BytesIO();document.save(output);return output.getvalue()
