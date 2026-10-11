"""Generate meeting-record Word documents from saved Nilaris transcripts."""

from datetime import datetime
from io import BytesIO
import re
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
import jdatetime


PERSIAN_MONTHS = (
    "فروردین", "اردیبهشت", "خرداد", "تیر",
    "مرداد", "شهریور", "مهر", "آبان",
    "آذر", "دی", "بهمن", "اسفند",
)

ENGLISH_MONTHS = (
    "January", "February", "March", "April",
    "May", "June", "July", "August",
    "September", "October", "November", "December",
)

SESSION_PATTERN = re.compile(
    r"^nilaris_\d{4}-\d{2}-\d{2}T\d{2}-\d{2}-\d{2}Z$"
)


def persian_digits(value):
    return str(value).translate(
        str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")
    )


def format_duration(seconds):
    total = max(0, int(round(float(seconds))))
    hours, remainder = divmod(total, 3600)
    minutes, seconds = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def set_text_direction(paragraph, rtl):
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
        if rtl
        else WD_ALIGN_PARAGRAPH.LEFT
    )

    properties = paragraph._p.get_or_add_pPr()

    bidi = properties.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        properties.append(bidi)

    bidi.set(qn("w:val"), "1" if rtl else "0")


def add_paragraph(document, text, *, rtl=False, bold=False, size=11):
    paragraph = document.add_paragraph()
    set_text_direction(paragraph, rtl)

    paragraph.paragraph_format.space_after = Pt(7)

    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)

    if rtl:
        properties = run._r.get_or_add_rPr()
        rtl_element = OxmlElement("w:rtl")
        rtl_element.set(qn("w:val"), "1")
        properties.append(rtl_element)

    return paragraph


def get_start_datetime(metadata, session_id):
    raw = metadata.get("start_time_utc")

    if raw:
        timestamp = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        if timestamp.tzinfo is None:
            raise ValueError("Session start time has no timezone.")
    else:
        if not SESSION_PATTERN.fullmatch(session_id):
            raise ValueError("Invalid session ID.")

        timestamp = datetime.strptime(
            session_id.removeprefix("nilaris_"),
            "%Y-%m-%dT%H-%M-%SZ",
        ).replace(tzinfo=ZoneInfo("UTC"))

    return timestamp.astimezone(ZoneInfo("Europe/Berlin"))


def format_dates(timestamp, language):
    gregorian = (
        f"{timestamp.day} "
        f"{ENGLISH_MONTHS[timestamp.month - 1]} "
        f"{timestamp.year}"
    )

    if language == "en":
        return [("Date", gregorian)]

    jalali = jdatetime.date.fromgregorian(
        date=timestamp.date()
    )

    jalali_text = (
        f"{persian_digits(jalali.day)} "
        f"{PERSIAN_MONTHS[jalali.month - 1]} "
        f"{persian_digits(jalali.year)}"
    )

    gregorian_fa = (
        f"{persian_digits(timestamp.day)} "
        f"{ENGLISH_MONTHS[timestamp.month - 1]} "
        f"{persian_digits(timestamp.year)}"
    )

    return [
        ("تاریخ شمسی", jalali_text),
        ("تاریخ میلادی", gregorian_fa),
    ]


def source_label(source, language):
    labels = {
        "LOCAL": {
            "fa": "صدای محیط جلسه",
            "en": "In-room audio",
        },
        "REMOTE": {
            "fa": "صدای شرکت‌کنندگان در تماس",
            "en": "Online participants' audio",
        },
    }

    return labels.get(
        str(source).upper(),
        {
            "fa": "صدای نامشخص",
            "en": "Unidentified audio",
        },
    )[language]


def segment_is_rtl(text):
    """Use RTL when the segment contains Persian/Arabic script."""
    return bool(re.search(r"[\u0600-\u06FF]", text))


def create_transcript_docx(transcript, metadata, language="fa"):
    """Return DOCX bytes without changing stored transcripts."""

    if language not in ("fa", "en"):
        raise ValueError("Language must be 'fa' or 'en'.")

    if transcript.get("timestamp_unit") != "seconds":
        raise ValueError("Unsupported transcript timestamp unit.")

    segments = transcript.get("segments")
    if not isinstance(segments, list):
        raise ValueError("Transcript segments must be a list.")

    session_id = metadata.get("session_id", "")
    if not SESSION_PATTERN.fullmatch(session_id):
        session_id = transcript.get("session_id", "").removesuffix(
            "_separated"
        )

    start_time = get_start_datetime(metadata, session_id)
    rtl = language == "fa"

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    title = "صورت‌جلسه" if rtl else "Meeting Record"
    subtitle = "متن کامل گفتگو" if rtl else "Full Meeting Transcript"

    add_paragraph(document, title, rtl=rtl, bold=True, size=17)
    add_paragraph(document, "NILARIS", rtl=rtl, bold=True, size=11)

    for label, value in format_dates(start_time, language):
        add_paragraph(document, f"{label}: {value}", rtl=rtl)

    time_value = start_time.strftime("%H:%M")
    duration_value = format_duration(metadata.get("duration_seconds", 0))

    if rtl:
        add_paragraph(
            document,
            f"ساعت شروع: {persian_digits(time_value)} به وقت آلمان",
            rtl=True,
        )
        add_paragraph(
            document,
            f"مدت ضبط: {persian_digits(duration_value)}",
            rtl=True,
        )
    else:
        add_paragraph(
            document,
            f"Start time: {time_value} (Berlin time)",
        )
        add_paragraph(document, f"Recording duration: {duration_value}")

    document.add_paragraph()
    add_paragraph(document, subtitle, rtl=rtl, bold=True, size=13)

    for segment in segments:
        start = format_duration(segment["start"])
        end = format_duration(segment["end"])
        label = source_label(segment.get("source"), language)

        if rtl:
            header = (
                f"{persian_digits(start)} – {persian_digits(end)}"
                f" | {label}"
            )
        else:
            header = f"{start} – {end} | {label}"

        add_paragraph(document, header, rtl=rtl, bold=True, size=10)

        # Preserve every segment's recognized text unchanged.
        add_paragraph(
            document,
            segment["text"],
            rtl=segment_is_rtl(segment["text"]),
            size=11,
        )

    output = BytesIO()
    document.save(output)
    return output.getvalue()
